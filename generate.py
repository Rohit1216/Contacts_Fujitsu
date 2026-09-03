"""
CLI usage:
    python generate.py template.pptx workbook.xlsx output.pptx

Sheet-to-slide mapping (edit SHEET_SLIDE_MAP below to match your template):
    workbook sheet name  ->  0-based index of the template slide to use
Any sheet not listed is skipped for the "team grid" pages, but every sheet
IS included (deduped) in the contact table.
"""

import sys

from pptx import Presentation

from utils.excel_loader import load_workbook_people, dedupe_by_linkedin
from utils.ppt_generator import (
    DeckTheme,
    paginate_team_sheet,
    paginate_contact_table,
)

# --- configure this per template -----------------------------------------
SHEET_SLIDE_MAP = {
    "Executive Management": 0,   # slide1.xml -> "Executive Management (x/y)"
    "IT": 1,                     # slide2.xml -> "IT Leadership Team (x/y)"
}
CONTACT_TABLE_SLIDE = 2          # slide3.xml -> "Contact Details..." table
TITLE_PREFIXES = {
    0: "Executive Management",
    1: "IT Leadership Team",
}
# ---------------------------------------------------------------------------


def main(template_path, workbook_path, output_path):
    prs = Presentation(template_path)
    theme = DeckTheme()
    photo_cache = {}

    people_by_sheet = load_workbook_people(workbook_path)

    # keep references to template slides by their ORIGINAL index before any
    # duplication shifts things around
    original_slides = list(prs.slides)
    template_slides = {idx: original_slides[idx] for idx in SHEET_SLIDE_MAP.values()}
    contact_template = original_slides[CONTACT_TABLE_SLIDE]

    for sheet_name, slide_idx in SHEET_SLIDE_MAP.items():
        people = people_by_sheet.get(sheet_name, [])
        print(f"[{sheet_name}] {len(people)} people -> slide index {slide_idx}")
        paginate_team_sheet(
            prs,
            template_slides[slide_idx],
            people,
            theme,
            photo_cache,
            title_prefix=TITLE_PREFIXES.get(slide_idx),
        )

    all_people = dedupe_by_linkedin(*people_by_sheet.values())
    print(f"[Contact table] {len(all_people)} unique people")
    paginate_contact_table(
        prs,
        contact_template,
        all_people,
        theme,
        title_prefix="Contact Details* of the Account Stakeholders",
    )

    prs.save(output_path)
    print("Saved:", output_path)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2], sys.argv[3])
